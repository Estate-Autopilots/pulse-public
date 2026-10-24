"""FIN-LIVE: live finance parsing, privacy, Drive confinement and the founders' brief. Synthetic data only."""

import io
import json
from datetime import UTC, date, datetime
from decimal import Decimal

import openpyxl
import pytest

from pulse_workers.finance_desk import synthetic_pdf
from pulse_workers.finance_live import bank, mail, model, sheets, sync
from pulse_workers.finance_live.common import as_date, money, tab_month
from pulse_workers.finance_live.drive import FINANCE_DRIVE, OutsideFinanceDrive, SharedDrive, file_name

PEOPLE = bank.People(founders=["Asha Founder"], staff=["Ravi Staffer", "Meena Editor"])


def workbook(tabs):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in tabs.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    data = io.BytesIO()
    wb.save(data)
    return openpyxl.load_workbook(io.BytesIO(data.getvalue()), data_only=True, read_only=True)


META = {"id": "synthetic-file", "name": "Synthetic sheet"}


def test_bank_truncated_names_match_the_invoice_client():
    from pulse_workers.finance_live.common import similar

    assert similar("Buvatur Commercial Development Company", "BUVATUR COMMERCIAL DEV COMPA") == 1.0
    assert similar("Sefavib Diamond Developers Corp", "SEFAVIB DIAMOND DEVELOPERS C") >= 0.75
    assert similar("Basil Avante LLP", "Basil Promoters LLP") < 0.8


def test_indian_amounts_dates_and_tab_months():
    assert money("1,10,000.00") == Decimal("110000.00")
    assert money("₹27,140.00") == Decimal("27140.00")
    assert money("INB/NEFT/123") is None
    assert as_date("09/10/2026") == date(2026, 10, 9)
    assert as_date("29-Feb-2024") == date(2024, 2, 29)
    assert tab_month("aug 25", None) == "2025-08"
    assert tab_month("Aprl", 2024) == "2024-04"
    assert tab_month("Trib & Investment", 2025) is None


def test_individual_payees_are_reduced_to_a_category_and_never_named():
    out = bank.classify("out", "INB/NEFT/AXOIC0000000001/Ravi Staffer/SYNTH BANK//////", PEOPLE)
    assert out["category"] == bank.PEOPLE and out["counterparty"] is None
    freelance = bank.classify("out", "INB/NEFT/AXOIC0000000002/Kedab Model/SYNTH BANK//", PEOPLE)
    assert freelance["category"] == bank.FREELANCE and freelance["counterparty"] is None
    vendor = bank.classify("out", "INB/NEFT/AXOIC0000000003/SYNTH SOFTWARE PRIVATE LIMITED/BANK/", PEOPLE)
    assert vendor["kind"] == "vendor" and vendor["counterparty"]
    founder_in = bank.classify("in", "IMPS/P2A/500000000001/ASHAFOUN/SYNTHBK/CAPITAL/919000", PEOPLE)
    assert founder_in["category"] == "Founder transfers in" and founder_in["counterparty"] is None
    client = bank.classify("in", "NEFT/HDFCH00000000001/SYNTH DEVELOPERS/HDFC BANK/Paid Inv 3033", PEOPLE)
    assert client["category"] == bank.CLIENT_RECEIPTS and client["invoices"] == ["INV-00003033"]


def test_cost_sheet_yields_payroll_total_and_cost_lines_without_names():
    wb = workbook(
        {
            "June 26": [
                [None, None, "Name", None, "Per Task", "Salary", "Travel", "total"],
                [None, "Video editor"],
                [1, None, "Ravi Staffer", 2, 500, 20000, 0, 21000],
                [2, None, "Meena Editor", 0, 0, 25000, 500, 25500],
                [None, "Model + Traveling", None, None, None, None, None, 9000],
                [None, "Fullfilment Total", None, None, None, None, None, 55500],
                [None, "Rent", None, None, None, None, None, 40000],
                [None, "Tools", None, None, None, None, None, 4000],
                [None, "Other Total", None, None, None, None, None, 44000],
                [None, "Founder Team"],
                [1, None, "Asha Founder", None, None, None, None, 50000],
                [None, "F.T. Total", None, None, None, None, None, 50000],
                [None, None, None, None, None, None, "Sub-Total", 149500],
            ]
        }
    )
    [month] = sheets.monthly_costs(wb, META, 2026)
    assert month["month"] == "2026-06" and month["people"] == 3
    assert month["payroll"] == Decimal("96500")
    assert month["items"] == {
        "Production & shoots": Decimal("9000"),
        "Office & rent": Decimal("40000"),
        "Software & AI": Decimal("4000"),
    }
    assert month["sheetTotal"] == month["payroll"] + sum(month["items"].values())
    dumped = json.dumps(month, default=str)
    assert "Ravi" not in dumped and "Meena" not in dumped and "Asha" not in dumped


def test_one_bank_line_copied_into_two_sheets_counts_once_and_statement_wins():
    lines = [
        {
            "date": "2025-08-01",
            "month": "2025-08",
            "direction": "in",
            "amount": Decimal("50000"),
            "text": "NEFT/HDFCH00000000009/SYNTH REALTY/HDFC BANK/",
            "family": "dated",
            "confidence": "medium",
            "source": {"type": "sheet", "fileId": "a", "file": "A", "sheet": "aug", "row": 2},
        },
        {
            "date": "2025-08-01",
            "month": "2025-08",
            "direction": "in",
            "amount": Decimal("50000"),
            "text": "NEFT/HDFCH00000000009/SYNTH REALTY/HDFC BANK/",
            "family": "dated",
            "confidence": "medium",
            "source": {"type": "sheet", "fileId": "b", "file": "B", "sheet": "aug 25", "row": 4},
        },
    ] + [
        {
            "date": f"2025-08-{d:02d}",
            "month": "2025-08",
            "direction": "in",
            "amount": Decimal("50000"),
            "text": f"NEFT/HDFCH0000000001{d}/SYNTH REALTY/HDFC BANK/",
            "family": "statement",
            "confidence": "high",
            "source": {"type": "sheet", "fileId": "s", "file": "S", "sheet": "Statement", "row": d},
        }
        for d in range(1, 13)
    ]
    result = sync.Result()
    sync.select_bank(result, lines, PEOPLE)
    rows = list(result.records["bank_line"].values())
    dated = [r for r in rows if r["body"]["family"] == "dated"]
    assert len(dated) == 1 and not dated[0]["body"]["selected"]
    assert all(r["body"]["selected"] for r in rows if r["body"]["family"] == "statement")


def test_finance_mail_selection_rejects_non_finance_and_otp_mail():
    assert (
        mail.rule_for(
            "marketing <message-service@sender.zohoinvoice.in>", "Invoice - INV-00003001 from Vikriya Solutions LLP"
        ).kind
        == "sales"
    )
    wifi = mail.rule_for("user-3357d85e@airtel.com", "Bill for your Airtel Wi-Fi 0000_wifi - June'26")
    assert wifi.vendor == "Airtel Wi-Fi"
    assert mail.rule_for("user-3357d85e@airtel.com", "Your OTP for login") is None
    assert mail.rule_for("friend@example.test", "Invoice for our dinner") is None
    assert mail.rule_for("Amazon Business <no-reply@business.amazon.in>", "Claim cashback on your purchases") is None
    group = mail.rule_for("new-vendor@example.test", "Your March statement", "billing@estateautopilots.com")
    assert group.vendor == "billing@ group"
    assert mail.rule_for("new-vendor@example.test", "Your OTP", "billing@estateautopilots.com") is None


def test_mail_attachments_are_fetched_only_for_selected_finance_messages():
    class Index:
        def __init__(self, rows):
            self.rows = rows

        def execute(self, sql, params=None):
            class R:
                def __init__(self, rows):
                    self._rows = rows

                def fetchall(self):
                    return self._rows

            return R(self.rows)

    sent = datetime(2026, 9, 1, 10, tzinfo=UTC)
    rows = [
        {
            "user_email": "marketing@estateautopilots.com",
            "id": "fin-1",
            "thread_id": "t",
            "from_addr": "user-3357d85e@airtel.com",
            "subject": "Bill for your Airtel Wi-Fi 1_wifi - August'26",
            "sent_at": sent,
            "snippet": "Amount due Rs 1,180.00",
            "size": 1,
        },
        {
            "user_email": "marketing@estateautopilots.com",
            "id": "private-1",
            "thread_id": "t",
            "from_addr": "client@example.test",
            "subject": "Campaign feedback",
            "sent_at": sent,
            "snippet": "private",
            "size": 1,
        },
    ]
    fetched = []
    original = sync._attachments
    sync._attachments = lambda gmail, message_id: fetched.append(message_id) or []
    try:
        result = sync.Result()
        sync.read_mail(result, Index(rows), lambda mailbox: object(), set(), {}, fetch=True)
    finally:
        sync._attachments = original
    assert fetched == ["fin-1"]
    [bill] = result.records["bill"].values()
    assert bill["body"]["total"] == 1180.0 and bill["body"]["basis"] == "mail-header"


def test_zoho_sales_invoice_pdf_is_parsed_with_tax_arithmetic():
    text = (
        "Sub Total 23,000.00\nCGST9 (9%) 2,070.00\nSGST9 (9%) 2,070.00\nTotal ₹27,140.00\nBalance Due ₹27,140.00\n"
        "GSTIN 27AAAAA0000A1Z5\nTAX INVOICE\n # : INV-00003001\n Invoice Date : 09/10/2026\n Due Date : 13/10/2026\n"
        " Place Of Supply : Maharashtra (27)\nSYNTH PROMOTERS LLP\nGSTIN 27BBBBB1111B1Z6\n"
    )
    parsed = mail.zoho_invoice(text, "27AAAAA0000A1Z5")
    assert parsed["number"] == "INV-00003001" and parsed["date"] == "2026-10-09" and parsed["due"] == "2026-10-13"
    assert parsed["client"] == "Synth Promoters LLP" and parsed["clientGstin"] == "27BBBBB1111B1Z6"
    assert parsed["total"] == Decimal("27140.00") and parsed["taxable"] == Decimal("23000.00") and not parsed["issues"]
    pdf = mail.pdf_text(synthetic_pdf("Invoice number: SYN-9\nAmount paid INR 1,180.00\nDate paid 01/09/2026"))
    assert mail.bill(pdf, "27AAAAA0000A1Z5")["total"] == Decimal("1180.00")


def test_gst_arn_mail_reveals_the_return_period():
    row = {
        "subject": "Acknowledgment (ARN) on filing of GSTR-3B",
        "snippet": "Acknowledgement Reference Number (ARN) AD270626000001X You have filed",
        "from_addr": "donotreply@gst.gov.in",
        "sent_at": datetime(2026, 7, 23, tzinfo=UTC),
    }
    f = mail.filing(row)
    assert f["type"] == "GSTR-3B" and f["period"] == "2026-06" and f["status"] == "filed"
    notice = {"from_addr": "PT Officer <officer@mahagst.gov.in>", "subject": "Show cause notice: Profession Tax"}
    pt = mail.filing({**row, **notice})
    assert pt["type"] == "Profession tax" and pt["status"] == "notice"


class FakeDrive:
    """A Shared Drive with one foreign folder. Records every call; supports only get/list/create."""

    def __init__(self):
        self.items = {
            FINANCE_DRIVE: {
                "id": FINANCE_DRIVE,
                "driveId": FINANCE_DRIVE,
                "mimeType": "application/vnd.google-apps.folder",
            },
            "foreign": {"id": "foreign", "driveId": "SOMEONE-ELSE", "mimeType": "application/vnd.google-apps.folder"},
        }
        self.calls = []
        self.n = 0

    def files(self):
        drive = self

        class Op:
            def __init__(self, fn):
                self.fn = fn

            def execute(self):
                return self.fn()

        class Files:
            def get(self, fileId, **_):
                drive.calls.append(("get", fileId))
                return Op(lambda: drive.items[fileId])

            def list(self, q, **_):
                drive.calls.append(("list", q))
                parent = q.split("'")[1]
                key = q.split("value='")[1].split("'")[0]
                return Op(
                    lambda: {
                        "files": [i for i in drive.items.values() if i.get("parent") == parent and i.get("key") == key]
                    }
                )

            def create(self, body, **_):
                drive.calls.append(("create", body["name"]))
                parent = body["parents"][0]
                drive.n += 1
                item = {
                    "id": f"new-{drive.n}",
                    "driveId": drive.items[parent]["driveId"],
                    "mimeType": body["mimeType"],
                    "parent": parent,
                    "key": body["appProperties"]["pulseFinanceKey"],
                }
                drive.items[item["id"]] = item
                return Op(lambda: item)

        return Files()


def test_drive_filing_is_confined_to_the_finance_shared_drive_and_idempotent():
    fake = FakeDrive()
    drive = SharedDrive(fake, FINANCE_DRIVE)
    pdf = synthetic_pdf("Synthetic invoice")
    name = file_name("2026-10-09", "Synth Promoters LLP", 27140, "INV-00003001")
    assert name == "2026-10-09_Synth_Promoters_LLP_27140.00_INV-00003001.pdf"
    first = drive.put(["FY2026-27", "Sales invoices", "2026-10"], name, pdf, "pdf-key")
    again = SharedDrive(fake, FINANCE_DRIVE).put(["FY2026-27", "Sales invoices", "2026-10"], name, pdf, "pdf-key")
    assert first["created"] and not again["created"] and first["id"] == again["id"]
    with pytest.raises(OutsideFinanceDrive):
        drive.folder("foreign", "FY2026-27")
    with pytest.raises(ValueError):
        drive.put(["FY2026-27"], "../escape.pdf", pdf, "other")
    assert {call[0] for call in fake.calls} <= {"get", "list", "create"}


def sample_records():
    invoices = [
        {
            "number": f"INV-0000{3000 + n}",
            "date": f"2025-0{7 + n % 3}-1{n % 9}",
            "due": f"2025-0{7 + n % 3}-2{n % 9}",
            "client": ["Alpha Realty", "Beta Homes", "Gamma Infra", "Delta Spaces"][n % 4],
            "taxable": 100000.0,
            "total": 118000.0,
            "balanceDue": 118000.0,
            "confidence": "high",
            "source": {"mailbox": "m", "message": str(n)},
        }
        for n in range(1, 9)
    ] + [
        {
            "number": "INV-00000002",
            "date": "2025-07-03",
            "due": "2025-07-03",
            "client": "Business Testing",
            "taxable": 1.0,
            "total": 1.18,
            "test": True,
            "confidence": "high",
            "source": {},
        },
        {
            "number": "INV-00003050",
            "date": "2026-09-20",
            "due": "2026-09-25",
            "client": "Beta Homes",
            "taxable": 50000.0,
            "total": 59000.0,
            "balanceDue": 59000.0,
            "confidence": "high",
            "source": {},
        },
    ]
    bank_lines = [
        {
            "key": "c1",
            "date": "2025-07-20",
            "month": "2025-07",
            "direction": "in",
            "amount": 118000.0,
            "category": bank.CLIENT_RECEIPTS,
            "counterparty": "BETA HOMES",
            "kind": "client",
            "invoices": [],
            "family": "statement",
            "confidence": "high",
            "selected": True,
        },
        {
            "key": "c2",
            "date": "2025-08-21",
            "month": "2025-08",
            "direction": "in",
            "amount": 116000.0,
            "category": bank.CLIENT_RECEIPTS,
            "counterparty": "GAMMA INFRA",
            "kind": "client",
            "invoices": ["INV-00003002"],
            "family": "statement",
            "confidence": "high",
            "selected": True,
        },
        {
            "key": "d1",
            "date": "2025-08-05",
            "month": "2025-08",
            "direction": "out",
            "amount": 300000.0,
            "category": bank.PEOPLE,
            "counterparty": None,
            "kind": "individual",
            "invoices": [],
            "family": "statement",
            "confidence": "high",
            "selected": True,
        },
    ]
    costs = [
        {
            "month": f"2026-0{m}",
            "payroll": 600000.0,
            "people": 20,
            "items": {"Office & rent": 40000.0},
            "confidence": "medium",
            "source": {"file": "Synthetic cost sheet"},
        }
        for m in range(4, 10)
    ]
    filings = [
        {"type": "GSTR-3B", "status": "filed", "date": "2026-07-23", "period": "2026-06", "arn": "AD2706260000001"},
        {"type": "GST registration", "status": "approved", "date": "2026-04-02"},
    ]
    balances = [{"from": "2025-10-01", "to": "2025-10-07", "opening": 100.0, "closing": 50000.0, "file": "Statement"}]
    return {
        "invoice": invoices,
        "bank_line": bank_lines,
        "cost_month": costs,
        "filing": filings,
        "bank_balance": balances,
        "package_cost": [{"package": "4 Month", "deliveryCost": 72750.0, "price": 100000.0}],
        "subscription": [{"name": "Claude", "monthly": 14250.0, "category": "Software & AI"}],
    }


def test_brief_is_short_specific_and_every_item_carries_confidence():
    m = model.build(sample_records(), date(2026, 10, 10))
    assert 1 <= len(m["brief"]) <= 7
    assert all(
        item["confidence"] in ("high", "medium", "low") and item["text"] and item["basis"] for item in m["brief"]
    )
    assert not m["certified"]
    paid = {i["number"]: i["status"] for i in m["collections"]["invoices"]}
    assert paid["INV-00003001"] == "paid" and paid["INV-00003002"] == "paid"
    assert all(i["number"] != "INV-00000002" for i in m["collections"]["invoices"]), (
        "test invoices never count as revenue"
    )
    assert m["collections"]["overdue"]["confirmed"]["count"] >= 1
    assert m["collections"]["overdue"]["unverified"]["count"] >= 1
    assert any(x["kind"] == "missing_invoices" for x in m["anomalies"])
    assert m["burn"]["runwayMonths"] is None and "days old" in m["burn"]["runwayNote"]
    assert {x["who"] for x in m["missing"]}, "every gap names who supplies it"
    payroll_text = json.dumps(m["payroll"])
    assert "Ravi" not in payroll_text and "Beta" not in payroll_text


def test_forecast_projects_only_recent_open_invoices():
    m = model.build(sample_records(), date(2026, 10, 10))
    note = next(a for a in m["forecast"]["assumptions"] if "last 60 days" in a)
    assert "(1," in note  # only INV-00003050 is recent; the 2025 invoices without evidence are not projected
    assert len(m["forecast"]["months"]) == 6 and m["forecast"]["confidence"] == "low"


def test_vendor_bill_tax_lines_balance_or_stay_unknown():
    google = (
        "Invoice\n₹ 286.06\n₹ 51.49\n₹ 337.55\nSubtotal in INR\nIntegrated GST (18%)\n"
        "Total in INR\n29 Feb 2024Invoice date\n"
    )
    parsed = mail.bill(google, "27AAAAA0000A1Z5")
    assert parsed["taxable"] == Decimal("286.06") and parsed["igst"] == Decimal("51.49")
    assert parsed["total"] == Decimal("337.55")
    lump = mail.bill("Taxes 18.00\nTotal Amount ` 118.00\n", "27AAAAA0000A1Z5")
    assert lump["gst"] == Decimal("18.00") and lump["taxable"] == Decimal("100.00")
    unbalanced = mail.bill("Sub Total 100.00\nIGST 50.00\nTotal 118.00\n", "27AAAAA0000A1Z5")
    assert "taxable" not in unbalanced and unbalanced["total"] == Decimal("118.00")

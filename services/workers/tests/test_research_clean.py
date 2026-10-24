import io
from datetime import timedelta

from pulse_workers.research_clean import creative, fold, locality_key, parse_ts, person_key, project_key
from pulse_workers.research_engine import visit_estimate
from pulse_workers.sheet_rows import tables, workbook_text


def test_normalisation_identity_and_locality():
    assert fold("  𝐒𝐚𝐦𝐚𝐫𝐭𝐡 &amp; Co  ") == "samarth & co"
    assert person_key("Mr. Samarth Putig Sir") == "samarth putig"
    assert project_key("Copy of New Synthetic Phase 2 Project Marathi Leads") == "synthetic ph2"
    assert locality_key("Poonawale Pune") == "punawale"
    assert creative("No Hook Marathi")[0] == "no_hook"
    assert creative("MV Hook 2 Hindi")[0] == "model_video"
    assert creative("Hook #1.2 English") == ("hook", "1.2", "english")


def test_dates_assume_utc_and_convert_to_ist():
    assert parse_ts("2026-09-01T22:00:00Z").isoformat() == "2026-09-02T03:30:00+05:30"
    assert parse_ts("2026-09-01T22:00").utcoffset() == timedelta(hours=5, minutes=30)
    assert parse_ts("2026-02-30T22:00Z") is None


def test_csv_transport_retains_empty_cells_pipes_and_multiline_notes():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(["created_time", "ad_name", "campaign_name", "platform", "notes"])
    ws.append(["2026-09-01T00:00Z", None, "Hook 1 | Pune", "fb", "call\n## notes\nmore"])
    buffer = io.BytesIO()
    wb.save(buffer)
    body = workbook_text(buffer.getvalue())
    name, rows, mode = list(tables(body))[0]
    assert mode == "csv-v1" and name == "Leads"
    assert rows[1] == ["2026-09-01T00:00Z", "", "Hook 1 | Pune", "fb", "call\n## notes\nmore"]
    assert list(tables("## old\na | b\nc | d\n"))[0][2] == "legacy-collapsed"


def test_small_samples_never_get_a_person_score():
    assert visit_estimate(1, 1, 0.1) == (None, None, None, "thin")
    value, low, high, sufficiency = visit_estimate(20, 40, 0.1)
    assert value == 0.3 and low < value < high and sufficiency == "enough"


def test_shared_index_rollout_changes_only_lead_sheet_transport():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["Name", "Salary"])
    ws.append(["Synthetic Person", None])
    lead = wb.create_sheet("Leads")
    lead.append(["created_time", "ad_name", "platform"])
    lead.append(["2026-09-01T00:00Z", None, "fb"])
    buffer = io.BytesIO()
    wb.save(buffer)
    people, leads = list(tables(workbook_text(buffer.getvalue(), csv_leads_only=True)))
    assert people[2] == "legacy-collapsed"
    assert leads[2] == "csv-v1" and leads[1][1][1] == ""
